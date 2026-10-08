"""견적서 생성 — 미리보기(HTML) / 엑셀(xlsx) / PDF.
외부 프로그램(LibreOffice 등) 설치 없이 순수 파이썬 라이브러리만 사용합니다.
세 가지 출력 모두 layout_spec() 이 만든 같은 구조를 그리므로 서로 모양이 어긋나지 않습니다.
"""
import html
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

COMPANY_NAME = "㈜ 애드리절트"
COMPANY_ADDR = "서울특별시 강남구 세곡동 129-7 애드리절트"
COMPANY_CEO = "이 승 민"
COMPANY_ACCOUNT = "853801-04-136993 국민은행 주식회사 애드리절트"

XLSX_FONT = "맑은 고딕"
GRAY_HEAD = "D9D9D9"      # 표 머리글 배경
GRAY_LABEL = "F2F2F2"     # 정보표 라벨 배경
GRAY_PRICE = "EFEFEF"     # 제안가 열 배경
FINAL_BG = "FFF6D9"       # 최종 합계 강조
LINE = "000000"

pdfmetrics.registerFont(UnicodeCIDFont("HYSMyeongJo-Medium"))  # 한글 내장 폰트 (폰트 파일 없이도 동작)
PDF_FONT = "HYSMyeongJo-Medium"
PDF_FONT_BOLD = PDF_FONT


def _register_korean_font():
    """시스템에 한글 TTF가 있으면 PDF에 직접 넣어서 어떤 뷰어에서도 글자가 깨지지 않게 합니다.
    없으면 위의 내장 CID 폰트를 그대로 씁니다 (설치는 필요 없음)."""
    from reportlab.pdfbase.ttfonts import TTFont
    candidates = [
        ("MalgunGothic", r"C:\Windows\Fonts\malgun.ttf", r"C:\Windows\Fonts\malgunbd.ttf"),   # Windows
        ("AppleGothic", "/System/Library/Fonts/Supplemental/AppleGothic.ttf", None),          # macOS
        ("NanumGothic", "/usr/share/fonts/truetype/nanum/NanumGothic.ttf",
         "/usr/share/fonts/truetype/nanum/NanumGothicBold.ttf"),                              # Linux
        ("NotoSansKR", "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc",
         "/usr/share/fonts/opentype/noto/NotoSansCJK-Bold.ttc"),
    ]
    for name, regular, bold in candidates:
        if not Path(regular).exists():
            continue
        try:
            kw = {"subfontIndex": 1} if regular.endswith(".ttc") else {}   # .ttc는 한국어 서브폰트 선택
            pdfmetrics.registerFont(TTFont(name, regular, **kw))
            bold_name = name
            if bold and Path(bold).exists():
                pdfmetrics.registerFont(TTFont(name + "-Bold", bold, **kw))
                bold_name = name + "-Bold"
            return name, bold_name
        except Exception:  # noqa: BLE001  폰트 파일이 깨졌거나 형식이 다르면 다음 후보로
            continue
    return PDF_FONT, PDF_FONT


PDF_FONT, PDF_FONT_BOLD = _register_korean_font()


# ══ 공통 계산 ═══════════════════════════════════════════════════
def _to_num(v) -> float:
    try:
        return float(str(v).replace(",", "").replace("₩", "").strip() or 0)
    except ValueError:
        return 0.0


def line_discount(item: dict) -> float:
    typ, val, price = item.get("할인유형", ""), _to_num(item.get("할인값", 0)), _to_num(item.get("단가", 0))
    if typ == "%":
        return min(price, price * val / 100)
    if typ == "원":
        return min(price, val)
    return 0.0


def line_final(item: dict) -> float:
    return _to_num(item.get("단가", 0)) - line_discount(item)


def calc_totals(items: list, overall_type: str = "없음", overall_value=0) -> dict:
    list_total = sum(_to_num(it.get("단가", 0)) for it in items)
    item_discount_total = sum(line_discount(it) for it in items)
    after_item = list_total - item_discount_total
    overall_value = _to_num(overall_value)
    overall_amount = min(after_item, after_item * overall_value / 100) if overall_type == "%" else \
        (min(after_item, overall_value) if overall_type == "원" else 0.0)
    after_all = after_item - overall_amount
    vat = after_all * 0.1
    return {"list_total": list_total, "item_discount_total": item_discount_total, "after_item": after_item,
            "overall_amount": overall_amount, "after_all": after_all, "vat": vat, "final": after_all + vat}


def _won(v) -> str:
    return f"\u20a9{v:,.0f}"


def _won_pdf(v) -> str:
    """PDF 내장 한글 폰트에서는 U+20A9(₩)가 빈칸으로 나와서 U+FFE6(￦)를 씁니다."""
    return f"\uffe6{v:,.0f}"


def _kdate(iso: str) -> str:
    """2026-09-07 → 2026년 09월 07일"""
    parts = (iso or "").split("-")
    return f"{parts[0]}년 {parts[1]}월 {parts[2]}일" if len(parts) == 3 else (iso or "")


def _qty_text(item: dict) -> str:
    """수량/기간 (+ 할인이 걸린 항목은 할인 표기를 덧붙임)"""
    base = str(item.get("수량기간", "") or "")
    typ, val = item.get("할인유형", ""), _to_num(item.get("할인값", 0))
    if typ == "%" and val:
        return f"{base} ({val:g}% 할인)".strip()
    if typ == "원" and val:
        return f"{base} (-{val:,.0f} 할인)".strip()
    return base


def layout_spec(header: dict, items: list, overall_type: str = "없음", overall_value=0) -> dict:
    """견적서 한 장의 구조. 채널이 같은 항목끼리 묶고, 제안가는 묶음 합계로 표시합니다."""
    items = [it for it in items if str(it.get("상품", "")).strip()]
    if not items:
        raise ValueError("견적에 담긴 항목이 없습니다.")

    groups, i = [], 0
    while i < len(items):
        ch = items[i].get("채널", "")
        j = i
        while j + 1 < len(items) and items[j + 1].get("채널", "") == ch:
            j += 1
        members = items[i:j + 1]
        groups.append({"채널": ch, "rows": members, "합계": sum(line_final(m) for m in members)})
        i = j + 1

    totals = calc_totals(items, overall_type, overall_value)
    left_info = [("브 랜 드", header.get("병원명", "")), ("참       조", header.get("참조", "")),
                 ("견 적 일 자", _kdate(header.get("견적일자", ""))),
                 ("견적유효기간", header.get("유효기간", "")), ("전       화", header.get("전화", ""))]
    right_info = [("견적번호", header.get("견적번호", "")), ("부     서", header.get("부서", "")),
                  ("담 당 자", header.get("담당자", "")), ("직     책", header.get("직책", "")),
                  ("전화번호", header.get("전화번호", ""))]

    total_rows = [("상품 정상가 합계", totals["list_total"])]
    if totals["item_discount_total"]:
        total_rows.append(("상품별 할인", -totals["item_discount_total"]))
    if totals["overall_amount"]:
        label = f"전체 할인 ({_to_num(overall_value):g}%)" if overall_type == "%" else "전체 할인"
        total_rows.append((label, -totals["overall_amount"]))
    if len(total_rows) > 1:
        total_rows.append(("공급가액", totals["after_all"]))
    total_rows.append(("부가세 (10%)", totals["vat"]))

    return {
        "title": f"{header.get('병원명', '')} 견적서",
        "to_line": f"{header.get('병원명', '')} {header.get('받는분', '원장님')} 귀중",
        "lead": "요청하신 상품건에 대하여 하기와 같이 견적합니다.",
        "left_info": left_info, "right_info": right_info,
        "groups": groups, "totals": totals, "total_rows": total_rows,
        "final_label": "최종 합계 (VAT 포함)", "final_value": totals["final"],
    }


# ══ 1) 미리보기 (HTML) ══════════════════════════════════════════
def build_quote_html(header: dict, items: list, overall_type: str = "없음", overall_value=0) -> str:
    s = layout_spec(header, items, overall_type, overall_value)
    e = html.escape

    left_rows = "".join(f"<tr><td class='lbl'>{e(ll)}</td><td class='val'>{e(str(lv))}</td></tr>"
                        for ll, lv in s["left_info"])
    right_rows = "".join(f"<tr><td class='lbl'>{e(rl)}</td><td class='val'>{e(str(rv))}</td></tr>"
                         for rl, rv in s["right_info"])

    body = ""
    for g in s["groups"]:
        n = len(g["rows"])
        for idx, it in enumerate(g["rows"]):
            body += "<tr>"
            if idx == 0:
                body += f"<td class='ch' rowspan='{n}'>{e(str(g['채널']))}</td>"
            body += (f"<td class='prod'>{e(str(it.get('상품', '')))}</td>"
                     f"<td class='desc'>{e(str(it.get('상세', '')))}</td>"
                     f"<td class='qty'>{e(_qty_text(it))}</td>")
            if idx == 0:
                body += f"<td class='price' rowspan='{n}'>{_won(g['합계'])}</td>"
            body += "</tr>"

    tot = "".join(f"<tr><td class='tl'>{e(l)}</td><td class='tv'>{_won(v)}</td></tr>" for l, v in s["total_rows"])
    tot += (f"<tr><td class='tl final'>{e(s['final_label'])}</td>"
            f"<td class='tv final'>{_won(s['final_value'])}</td></tr>")

    return f"""<style>
#qpaper {{ background:#eef1f5; padding:14px; font-family:"Malgun Gothic","맑은 고딕",system-ui,sans-serif;
          color:#111; font-size:13px; }}
#qpaper * {{ box-sizing:border-box; }}
#qpaper .page {{ background:#fff; max-width:900px; margin:0 auto; padding:0 0 26px; border:1px solid #333; }}
#qpaper .titlebar {{ display:flex; justify-content:space-between; align-items:center;
                    padding:10px 16px; border-bottom:1px solid #333; }}
#qpaper .titlebar .t {{ font-size:17px; font-weight:700; }}
#qpaper .titlebar .logo {{ font-size:13px; font-weight:700; }}
#qpaper .inner {{ margin:16px; border:1px solid #333; padding:20px 24px 26px; }}
#qpaper .head {{ display:flex; justify-content:space-between; margin-bottom:16px; gap:20px; }}
#qpaper .head .r {{ text-align:center; }}
#qpaper .head .r .cn {{ font-weight:700; }}
#qpaper .ceo {{ display:flex; justify-content:space-between; gap:40px; margin-top:4px; letter-spacing:2px; }}
#qpaper table.info {{ border-collapse:collapse; margin-bottom:20px; }}
#qpaper table.info td {{ border:1px solid #333; padding:5px 10px; height:26px; font-size:12px; }}
#qpaper table.info td.lbl {{ background:#{GRAY_LABEL}; text-align:center; letter-spacing:1px;
                            white-space:nowrap; }}
#qpaper table.info td.val {{ text-align:center; min-width:190px; }}
#qpaper .inforow {{ display:flex; justify-content:space-between; gap:24px; margin-bottom:20px; }}
#qpaper table.items {{ border-collapse:collapse; width:100%; }}
#qpaper table.items th {{ background:#{GRAY_HEAD}; border:1px solid #333; padding:10px 6px;
                         font-size:13px; font-weight:700; }}
#qpaper table.items td {{ border:1px solid #333; padding:12px 8px; text-align:center;
                         vertical-align:middle; font-size:12px; }}
#qpaper table.items td.ch {{ font-weight:700; }}
#qpaper table.items td.desc {{ text-align:left; line-height:1.5; }}
#qpaper table.items td.price {{ background:#{GRAY_PRICE}; font-weight:700; white-space:nowrap; }}
#qpaper table.tot {{ border-collapse:collapse; width:100%; margin-top:16px; }}
#qpaper table.tot td {{ border:1px solid #333; padding:7px 14px; font-size:12px; }}
#qpaper table.tot td.tl {{ background:#{GRAY_LABEL}; font-weight:700; }}
#qpaper table.tot td.tv {{ text-align:right; width:135px; white-space:nowrap; }}
#qpaper table.tot td.final {{ font-size:14px; font-weight:700; background:#{FINAL_BG}; }}
#qpaper .acct {{ text-align:center; margin-top:24px; font-size:15px; font-weight:700; letter-spacing:.5px; }}
</style>
<div id="qpaper"><div class="page">
  <div class="titlebar"><div class="t">{e(s['title'])}</div><div class="logo">{e(COMPANY_NAME)}</div></div>
  <div class="inner">
    <div class="head">
      <div class="l"><div>{e(s['to_line'])}</div><div>{e(s['lead'])}</div></div>
      <div class="r"><div class="cn">{e(COMPANY_NAME)}</div><div>{e(COMPANY_ADDR)}</div>
        <div class="ceo"><span>대 표 이 사</span><span>{e(COMPANY_CEO)}</span></div></div>
    </div>
    <div class="inforow">
      <table class="info">{left_rows}</table>
      <table class="info">{right_rows}</table>
    </div>
    <table class="items">
      <tr><th style="width:115px">채널</th><th style="width:145px">상품</th><th>상세</th>
          <th style="width:135px">수량/기간</th><th style="width:135px">제안가</th></tr>
      {body}
    </table>
    <table class="tot">{tot}</table>
    <div class="acct">{e(COMPANY_ACCOUNT)}</div>
  </div>
</div></div>"""


# ══ 2) 엑셀 (xlsx) ══════════════════════════════════════════════
def _s():
    return Side(style="thin", color=LINE)


def _put(ws, coord, value, *, size=10, bold=False, align="center", fill=None, wrap=False, fmt=None,
         border=True):
    c = ws[coord]
    c.value = value
    c.font = Font(name=XLSX_FONT, size=size, bold=bold)
    c.alignment = Alignment(horizontal=align, vertical="center", wrap_text=wrap)
    if border:
        s = _s()
        c.border = Border(left=s, right=s, top=s, bottom=s)
    if fill:
        c.fill = PatternFill("solid", fgColor=fill)
    if fmt:
        c.number_format = fmt
    return c


def _region(ws, ref, *, fill=None, border=True):
    """병합 영역 전체에 테두리/배경을 적용 (병합해도 선이 끊기지 않게)"""
    s = _s()
    for row in ws[ref]:
        for c in row:
            if border:
                c.border = Border(left=s, right=s, top=s, bottom=s)
            if fill:
                c.fill = PatternFill("solid", fgColor=fill)


def build_quote_xlsx(out_path: str, header: dict, items: list, overall_type: str = "없음", overall_value=0) -> str:
    spec = layout_spec(header, items, overall_type, overall_value)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "견적서"
    ws.sheet_view.showGridLines = False

    for col, w in {"A": 3, "B": 16, "C": 22, "D": 13, "E": 13, "F": 13, "G": 8, "H": 15, "I": 10,
                   "J": 10, "K": 3}.items():
        ws.column_dimensions[col].width = w

    ws.merge_cells("B2:G2")
    _put(ws, "B2", spec["title"], size=15, bold=True, align="left", border=False)
    ws.merge_cells("H2:J2")
    _put(ws, "H2", COMPANY_NAME, size=11, bold=True, align="right", border=False)
    ws.row_dimensions[2].height = 26

    _put(ws, "B4", spec["to_line"], size=10, align="left", border=False)
    _put(ws, "B5", spec["lead"], size=9, align="left", border=False)
    ws.merge_cells("H4:J4")
    _put(ws, "H4", COMPANY_NAME, size=10, bold=True, border=False)
    ws.merge_cells("H5:J5")
    _put(ws, "H5", COMPANY_ADDR, size=9, border=False)
    ws.merge_cells("H6:I6")
    _put(ws, "H6", "대 표 이 사", size=9, border=False)
    _put(ws, "J6", COMPANY_CEO, size=9, border=False)

    r0 = 8
    for i, ((ll, lv), (rl, rv)) in enumerate(zip(spec["left_info"], spec["right_info"])):
        r = r0 + i
        _put(ws, f"B{r}", ll, size=9, fill=GRAY_LABEL)
        ws.merge_cells(f"C{r}:E{r}")
        _put(ws, f"C{r}", str(lv), size=9)
        _region(ws, f"C{r}:E{r}")
        _put(ws, f"G{r}", rl, size=9, fill=GRAY_LABEL)
        ws.merge_cells(f"H{r}:J{r}")
        _put(ws, f"H{r}", str(rv), size=9)
        _region(ws, f"H{r}:J{r}")
        ws.row_dimensions[r].height = 20

    head = r0 + len(spec["left_info"]) + 2
    _put(ws, f"B{head}", "채널", bold=True, fill=GRAY_HEAD)
    _put(ws, f"C{head}", "상품", bold=True, fill=GRAY_HEAD)
    ws.merge_cells(f"D{head}:G{head}")
    _put(ws, f"D{head}", "상세", bold=True, fill=GRAY_HEAD)
    _region(ws, f"D{head}:G{head}", fill=GRAY_HEAD)
    _put(ws, f"H{head}", "수량/기간", bold=True, fill=GRAY_HEAD)
    ws.merge_cells(f"I{head}:J{head}")
    _put(ws, f"I{head}", "제안가", bold=True, fill=GRAY_HEAD)
    _region(ws, f"I{head}:J{head}", fill=GRAY_HEAD)
    ws.row_dimensions[head].height = 24

    r = head + 1
    for g in spec["groups"]:
        start = r
        for it in g["rows"]:
            _put(ws, f"C{r}", str(it.get("상품", "")), size=9)
            ws.merge_cells(f"D{r}:G{r}")
            _put(ws, f"D{r}", str(it.get("상세", "")), size=9, align="left", wrap=True)
            _region(ws, f"D{r}:G{r}")
            _put(ws, f"H{r}", _qty_text(it), size=9, wrap=True)
            ws.row_dimensions[r].height = 40
            r += 1
        end = r - 1
        if end > start:
            ws.merge_cells(f"B{start}:B{end}")
        ws.merge_cells(f"I{start}:J{end}")
        _put(ws, f"B{start}", str(g["채널"]), size=10, bold=True)
        _region(ws, f"B{start}:B{end}")
        _put(ws, f"I{start}", g["합계"], size=10, bold=True, fill=GRAY_PRICE, fmt='₩#,##0')
        _region(ws, f"I{start}:J{end}", fill=GRAY_PRICE)

    r += 1
    for label, value in spec["total_rows"]:
        ws.merge_cells(f"B{r}:H{r}")
        _put(ws, f"B{r}", label, size=10, bold=True, align="left", fill=GRAY_LABEL)
        _region(ws, f"B{r}:H{r}", fill=GRAY_LABEL)
        ws.merge_cells(f"I{r}:J{r}")
        _put(ws, f"I{r}", value, size=10, align="right", fmt='₩#,##0')
        _region(ws, f"I{r}:J{r}")
        ws.row_dimensions[r].height = 20
        r += 1
    ws.merge_cells(f"B{r}:H{r}")
    _put(ws, f"B{r}", spec["final_label"], size=11, bold=True, align="left", fill=FINAL_BG)
    _region(ws, f"B{r}:H{r}", fill=FINAL_BG)
    ws.merge_cells(f"I{r}:J{r}")
    _put(ws, f"I{r}", spec["final_value"], size=12, bold=True, align="right", fill=FINAL_BG, fmt='₩#,##0')
    _region(ws, f"I{r}:J{r}", fill=FINAL_BG)
    ws.row_dimensions[r].height = 24

    acct = r + 2
    ws.merge_cells(f"B{acct}:J{acct}")
    _put(ws, f"B{acct}", COMPANY_ACCOUNT, size=13, bold=True, border=False)
    ws.row_dimensions[acct].height = 26

    ws.print_area = f"A1:K{acct + 1}"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    Path(out_path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return out_path


# ══ 3) PDF ══════════════════════════════════════════════════════
def _p(text, size=9, bold=False, align=1, leading=None):
    style = ParagraphStyle("s", fontName=PDF_FONT_BOLD if bold else PDF_FONT, fontSize=size,
                           leading=leading or size * 1.45, alignment=align)
    return Paragraph(str(text).replace("\n", "<br/>"), style)


def build_quote_pdf(out_path: str, header: dict, items: list, overall_type: str = "없음", overall_value=0) -> str:
    spec = layout_spec(header, items, overall_type, overall_value)
    Path(out_path).parent.mkdir(parents=True, exist_ok=True)
    doc = SimpleDocTemplate(out_path, pagesize=A4, topMargin=14 * mm, bottomMargin=14 * mm,
                            leftMargin=13 * mm, rightMargin=13 * mm, title=spec["title"])
    W = 184 * mm
    grid = colors.HexColor("#333333")
    story = []

    tb = Table([[_p(spec["title"], 14, align=0), _p(COMPANY_NAME, 10, align=2)]],
               colWidths=[W * 0.7, W * 0.3])
    tb.setStyle(TableStyle([("BOX", (0, 0), (-1, -1), 0.8, grid), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                            ("TOPPADDING", (0, 0), (-1, -1), 6), ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
                            ("LEFTPADDING", (0, 0), (0, 0), 8), ("RIGHTPADDING", (-1, -1), (-1, -1), 8)]))
    story += [tb, Spacer(1, 7 * mm)]

    ceo = Table([[_p("대 표 이 사", 8.5), _p(COMPANY_CEO, 8.5)]], colWidths=[W * 0.22, W * 0.22])
    head_tbl = Table([[_p(spec["to_line"], 9.5, align=0), _p(COMPANY_NAME, 9.5)],
                      [_p(spec["lead"], 8.5, align=0), _p(COMPANY_ADDR, 8.5)],
                      ["", ceo]], colWidths=[W * 0.52, W * 0.48])
    head_tbl.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("BOTTOMPADDING", (0, 0), (-1, -1), 3)]))
    story += [head_tbl, Spacer(1, 6 * mm)]

    lw, vw = W * 0.13, W * 0.29
    gap = W - (lw + vw) * 2          # 좌우 표를 양 끝으로 밀어내는 가운데 여백
    info_rows = [[_p(ll, 8.5), _p(str(lv), 8.5), "", _p(rl, 8.5), _p(str(rv), 8.5)]
                 for (ll, lv), (rl, rv) in zip(spec["left_info"], spec["right_info"])]
    info = Table(info_rows, colWidths=[lw, vw, gap, lw, vw])
    info.setStyle(TableStyle([
        ("GRID", (0, 0), (1, -1), 0.6, grid), ("GRID", (3, 0), (4, -1), 0.6, grid),
        ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#" + GRAY_LABEL)),
        ("BACKGROUND", (3, 0), (3, -1), colors.HexColor("#" + GRAY_LABEL)),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4)]))
    story += [info, Spacer(1, 8 * mm)]

    cw = [W * 0.15, W * 0.19, W * 0.34, W * 0.16, W * 0.16]
    rows = [[_p("채널", 9.5, True), _p("상품", 9.5, True), _p("상세", 9.5, True),
             _p("수량/기간", 9.5, True), _p("제안가", 9.5, True)]]
    style = [("GRID", (0, 0), (-1, -1), 0.6, grid),
             ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#" + GRAY_HEAD)),
             ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
             ("TOPPADDING", (0, 0), (-1, -1), 7), ("BOTTOMPADDING", (0, 0), (-1, -1), 7)]
    r = 1
    for g in spec["groups"]:
        start = r
        for k, it in enumerate(g["rows"]):
            rows.append([_p(g["채널"], 9, True) if k == 0 else "",
                         _p(it.get("상품", ""), 8.5),
                         _p(it.get("상세", ""), 8, align=0),
                         _p(_qty_text(it), 8.5),
                         _p(_won_pdf(g["합계"]), 9, True) if k == 0 else ""])
            r += 1
        end = r - 1
        if end > start:
            style += [("SPAN", (0, start), (0, end)), ("SPAN", (4, start), (4, end))]
        style.append(("BACKGROUND", (4, start), (4, end), colors.HexColor("#" + GRAY_PRICE)))
    item_tbl = Table(rows, colWidths=cw, repeatRows=1)
    item_tbl.setStyle(TableStyle(style))
    story += [item_tbl, Spacer(1, 6 * mm)]

    trows = [[_p(l, 9, True, align=0), _p(_won_pdf(v), 9, align=2)] for l, v in spec["total_rows"]]
    trows.append([_p(spec["final_label"], 10.5, True, align=0), _p(_won_pdf(spec["final_value"]), 11, True, align=2)])
    tot = Table(trows, colWidths=[W - W * 0.16, W * 0.16], hAlign="LEFT")   # 항목표와 같은 전체 폭
    tot.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.6, grid),
                             ("BACKGROUND", (0, 0), (0, -2), colors.HexColor("#" + GRAY_LABEL)),
                             ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#" + FINAL_BG)),
                             ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                             ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5)]))
    story += [tot, Spacer(1, 10 * mm), _p(COMPANY_ACCOUNT, 12.5, True)]

    doc.build(story)
    return out_path
